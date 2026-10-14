"""Regenerate the synthetic Office fixtures used by the preview tests.

Every value below is invented for testing; no user content is involved.

    python3 -m venv <TMP>/pyenv
    <TMP>/pyenv/bin/pip install python-docx==1.2.0 python-pptx==1.0.2 openpyxl==3.1.5
    <TMP>/pyenv/bin/python test/fixtures/office/generate.py
    soffice --headless -env:UserInstallation=file://<TMP>/lo-profile \
      --convert-to xlsx --outdir <TMP>/lo-out test/fixtures/office/sample-source.xlsx
    soffice --headless -env:UserInstallation=file://<TMP>/lo-profile \
      --convert-to xls --outdir <TMP>/lo-out test/fixtures/office/sample-source.xlsx

The LibreOffice round trip stores calculated values for formula cells, which
is what a workbook saved by a spreadsheet application contains. Copy the
converted files to sample.xlsx and sample.xls and delete sample-source.xlsx.
"""

import datetime
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as PptxInches, Pt

HERE = Path(__file__).resolve().parent


def png_bytes(width=48, height=32):
    """A small gradient PNG built without an imaging library."""
    rows = b"".join(
        b"\x00" + b"".join(bytes((x * 5 % 256, y * 7 % 256, 160)) for x in range(width))
        for y in range(height)
    )

    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    rel_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rel_id}">'
        f'<w:r><w:rPr><w:u w:val="single"/></w:rPr><w:t xml:space="preserve">{text}</w:t></w:r></w:hyperlink>'
    )
    paragraph._p.append(hyperlink)


def build_docx():
    document = Document()
    document.core_properties.author = "Fixture"
    document.add_heading("季度阅读报告 Quarterly reading report", level=0)
    document.add_heading("一、概述 Overview", level=1)
    intro = document.add_paragraph("本文件是测试用的合成文档。")
    intro.add_run("粗体 bold").bold = True
    intro.add_run("，")
    intro.add_run("斜体 italic").italic = True
    intro.add_run("，")
    intro.add_run("下划线 underline").underline = True
    intro.add_run("。")
    document.add_heading("二、清单 Lists", level=2)
    for item in ["第一项 first", "第二项 second", "第三项 third"]:
        document.add_paragraph(item, style="List Bullet")
    for item in ["步骤一 step one", "步骤二 step two"]:
        document.add_paragraph(item, style="List Number")
    document.add_heading("三、表格 Table", level=2)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row_index, row in enumerate([["姓名 Name", "篇数 Count", "备注 Note"], ["甲", "12", "按时"], ["乙", "9", "补交"]]):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    link_paragraph = document.add_paragraph("参考链接 Reference: ")
    add_hyperlink(link_paragraph, "https://example.org/reading", "example.org/reading")
    document.add_heading("四、图片 Image", level=2)
    image_path = HERE / "fixture-image.png"
    image_path.write_bytes(png_bytes())
    document.add_picture(str(image_path), width=Inches(1.5))
    image_path.unlink()
    document.add_heading("五、公式 Equation", level=2)
    equation = document.add_paragraph()
    equation._p.append(
        parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            "<m:r><m:t>E=mc</m:t></m:r><m:sSup><m:e><m:r><m:t> </m:t></m:r></m:e>"
            "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>"
        )
    )
    document.add_paragraph("文末 End of document.")
    document.save(HERE / "sample.docx")


def build_xlsx_source():
    workbook = Workbook()
    grades = workbook.active
    grades.title = "成绩 Scores"
    grades["A1"] = "第一学期阅读统计"
    grades["A1"].font = Font(bold=True)
    grades.merge_cells("A1:D1")
    grades.append(["姓名 Name", "篇数 Count", "比例 Ratio", "日期 Date"])
    rows = [("甲", 12, 0.4), ("乙", 9, 0.3), ("丙", 9, 0.3)]
    for index, (name, count, ratio) in enumerate(rows, start=3):
        grades.append([name, count, ratio, datetime.date(2026, 9, index)])
        grades[f"C{index}"].number_format = "0%"
        grades[f"D{index}"].number_format = "yyyy-mm-dd"
    grades["A6"] = "合计 Total"
    grades["B6"] = "=SUM(B3:B5)"
    grades["C6"] = "=B6/30"
    grades["C6"].number_format = "0.0%"
    grades.column_dimensions["A"].width = 18

    hidden = workbook.create_sheet("隐藏 Hidden")
    hidden["A1"] = "hidden sheet value"
    hidden.sheet_state = "hidden"

    long_sheet = workbook.create_sheet("长表 Long")
    long_sheet.append(["序号 Row", "值 Value", "说明 Note"])
    for index in range(1, 2501):
        long_sheet.append([index, index * 3, f"第 {index} 行"])

    wide_sheet = workbook.create_sheet("宽表 Wide")
    wide_sheet.append([f"C{index}" for index in range(1, 261)])

    formulas = workbook.create_sheet("公式 Formulas")
    formulas["A1"] = 2
    formulas["A2"] = 3
    formulas["A3"] = "=A1*A2"
    formulas["A4"] = '=HYPERLINK("https://example.org/","link")'
    formulas["A5"] = "=1/0"

    workbook.active = 0
    workbook.save(HERE / "sample-source.xlsx")


def build_pptx():
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "课程汇报 Course review"
    title_slide.placeholders[1].text = "合成测试演示文稿 Synthetic test deck"

    bullets = presentation.slides.add_slide(presentation.slide_layouts[1])
    bullets.shapes.title.text = "要点 Key points"
    body = bullets.placeholders[1].text_frame
    body.text = "阅读量提升 Reading volume up"
    for text in ["讨论更积极 More discussion", "作业按时 On-time homework"]:
        paragraph = body.add_paragraph()
        paragraph.text = text
        paragraph.level = 1
    link_run = body.add_paragraph().add_run()
    link_run.text = "示例链接 example link"
    link_run.hyperlink.address = "https://example.org/slides"

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "数据表 Table"
    table = table_slide.shapes.add_table(3, 3, PptxInches(1), PptxInches(2), PptxInches(8), PptxInches(1.5)).table
    for row_index, row in enumerate([["班级", "人数", "均分"], ["一班", "40", "86.5"], ["二班", "38", "84.0"]]):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    chart_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    chart_slide.shapes.title.text = "图表 Chart"
    chart_data = CategoryChartData()
    chart_data.categories = ["九月", "十月", "十一月"]
    chart_data.add_series("篇数 Count", (12, 18, 15))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, PptxInches(1), PptxInches(1.8), PptxInches(8), PptxInches(4.5), chart_data
    )

    picture_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    picture_slide.shapes.title.text = "图片 Picture"
    image_path = HERE / "fixture-image.png"
    image_path.write_bytes(png_bytes())
    picture_slide.shapes.add_picture(str(image_path), PptxInches(3), PptxInches(2.5), width=PptxInches(4))
    image_path.unlink()
    note = picture_slide.shapes.add_textbox(PptxInches(1), PptxInches(6.2), PptxInches(8), PptxInches(0.6))
    note.text_frame.text = "动画、音频和视频不在预览中播放。"
    note.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    presentation.save(HERE / "sample.pptx")


if __name__ == "__main__":
    build_docx()
    build_xlsx_source()
    build_pptx()
